#!/usr/bin/env python3
"""Builds the working report from a FINAL_FILTERED export.

The export is a data file: every field the filter computed, one row per
organisation. What people actually work with is narrower and sorted - which is
what this produces: a summary, one sheet per district, the religious
institutions kept apart from food, the chains counted, and the party branches
from a hand-kept file appended as their own sheet.

    pip install openpyxl
    python3 tools/report.py --final FINAL_FILTERED-*.csv --out REESTR.xlsx
    python3 tools/report.py --final … --er "ЕР.xlsx" --per-district out/po_rayonam/
    python3 tools/report.py --finalize --out REESTR.xlsx --per-district out/po_rayonam/

--finalize takes books that already exist (possibly edited by hand) and only
puts them in order: a Справка sheet first in every book, the party branches and
the religious sheet right after the summary, districts numbered 01-12 in the
same order as the per-district files. build() runs the same step at the end.

Rules that matter and must not drift:
  * an organisation is identified by place_id, never by name or distance;
  * its district is detected_district (from coordinates), not the query's;
  * a row whose groups include Религиозные учреждения goes to the religious
    sheet only - never into food, whatever else Yandex filed it under.
"""
import argparse, csv, datetime, glob, io, math, re, collections, os, sys

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit('нужен openpyxl:  pip install openpyxl')

csv.field_size_limit(10 ** 7)

RELIGION = 'Религиозные учреждения'
GROUP_ORDER = {'Общепит': 0, 'Продуктовая розница': 1, 'Пищевое производство': 2, RELIGION: 3}
HEAD = Font(bold=True, color='FFFFFF', size=11)
FILL = PatternFill('solid', fgColor='2F4858')
THIN = Border(bottom=Side(style='thin', color='DDDDDD'))
LINK = Font(color='1155CC', underline='single')
NOTE = Font(italic=True, size=9, color='555555')

COLS = [('№', 6), ('Название', 34), ('Группа', 24), ('Рубрики', 32), ('Адрес', 38), ('Телефон', 19),
        ('Доб.', 7), ('Сайт', 30), ('Часы работы', 26), ('Рейтинг', 8), ('Отзывов', 9), ('Сеть', 7),
        ('Точек в сети', 12), ('По адресу', 10), ('Карточка', 11), ('place_id', 14)]
WRAP = {'Рубрики', 'Адрес', 'Часы работы', 'Возможный дубль'}


def read_final(path):
    with io.open(path, encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    if not rows:
        sys.exit(f'{path}: пустой файл')
    missing = {'detected_district', 'detected_category_group', 'place_id'} - set(rows[0])
    if missing:
        sys.exit(f'{path}: это не FINAL_FILTERED — нет колонок {", ".join(sorted(missing))}')
    return rows


def groups_of(row):
    found = [g.strip() for g in (row.get('detected_category_group') or '').split('|') if g.strip()]
    found.sort(key=lambda g: GROUP_ORDER.get(g, 9))
    return found


def label(row):
    return ' / '.join(groups_of(row)) or '—'


def norm_addr(value):
    text = (value or '').lower().replace('ё', 'е')
    text = re.sub(r'\s*,\s*москва\s*$', '', text)
    return re.sub(r'\s+', ' ', text).strip()


def metres(a, b):
    try:
        lat_a, lon_a = float(a['latitude']), float(a['longitude'])
        lat_b, lon_b = float(b['latitude']), float(b['longitude'])
    except (TypeError, ValueError):
        return float('inf')
    return math.hypot((lat_a - lat_b) * 111320, (lon_a - lon_b) * 111320 * math.cos(math.radians(lat_a)))


def first_rubric(row):
    return (row.get('categories') or '').split(',')[0].strip().lower().replace('ё', 'е')


def annotate(rows, near_metres=40):
    """Counts organisations per address, and names the ones standing on top of
    each other. Both are hints for a human: nothing is merged or dropped here,
    because two churches in one yard are two churches."""
    per_address = collections.Counter(norm_addr(r['address']) for r in rows if norm_addr(r['address']))
    near = collections.defaultdict(list)
    for i in range(len(rows)):
        for j in range(i + 1, len(rows)):
            if first_rubric(rows[i]) != first_rubric(rows[j]):
                continue
            gap = metres(rows[i], rows[j])
            if gap <= near_metres:
                near[i].append((round(gap), rows[j]['title']))
                near[j].append((round(gap), rows[i]['title']))
    for i, row in enumerate(rows):
        row['_per_addr'] = per_address.get(norm_addr(row['address']), 0)
        row['_dup'] = '; '.join(f'{t} ({d} м)' for d, t in sorted(near.get(i, []))[:3])
    return rows


def number(value, cast=float):
    try:
        return cast(value)
    except (TypeError, ValueError):
        return ''


def working_sheet(wb, name, rows, district_col=False, dup=False):
    sheet = wb.create_sheet(name[:31])
    cols = list(COLS)
    if district_col:
        cols = [cols[0], ('Район', 18)] + cols[1:]
    if dup:
        cols = cols + [('Возможный дубль', 46)]
    names = [c for c, _ in cols]
    sheet.append(names)
    for n, row in enumerate(rows, 1):
        line = [n]
        if district_col:
            line.append(row['detected_district'])
        line += [row['title'], label(row), row['categories'], row['address'], row['phone'],
                 row.get('phone_ext', ''), row['website'], row['opening_hours'],
                 number(row.get('rating')), number(row.get('review_count'), int),
                 'Да' if row.get('chain_flag') == 'CHAIN' else '', number(row.get('chain_size'), int),
                 row['_per_addr'] if row['_per_addr'] > 1 else '',
                 'открыть' if row['maps_url'] else '', row['place_id']]
        if dup:
            line.append(row['_dup'])
        sheet.append(line)
        i = sheet.max_row
        for column, url in (('Сайт', row['website']), ('Карточка', row['maps_url'])):
            if url:
                cell = sheet.cell(i, names.index(column) + 1)
                cell.hyperlink = url
                cell.font = LINK
        for c in range(1, len(cols) + 1):
            sheet.cell(i, c).border = THIN
    for cell in sheet[1]:
        cell.font, cell.fill = HEAD, FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 30
    for i, (_, width) in enumerate(cols, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for column in WRAP & set(names):
        letter = get_column_letter(names.index(column) + 1)
        for i in range(2, sheet.max_row + 1):
            sheet[f'{letter}{i}'].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{sheet.max_row}'
    return sheet


def er_sheets(path):
    """The party branches are kept by hand, not collected: taken as they are."""
    wb = load_workbook(path)
    header, rows = None, []
    for name in wb.sheetnames:
        for raw in wb[name].iter_rows(values_only=True):
            cells = ['' if x is None else str(x) for x in raw]
            if not any(cells):
                continue
            if cells[0].strip() == 'Район':
                header = header or cells
                continue
            if header and len(cells) >= 2:
                rows.append(cells)
    return header, rows


FOOD_GROUPS = ('Общепит', 'Продуктовая розница', 'Пищевое производство')
FIXED_SHEETS = ('Справка', 'Сводка', 'Единая Россия', 'Религия')
NUMBERED = re.compile(r'^\d\d ')


def plain(name):
    return NUMBERED.sub('', name)


def sheet_rows(sheet):
    """Header and non-empty data rows of a working sheet."""
    rows = [r for r in sheet.iter_rows(values_only=True) if any(v not in (None, '') for v in r)]
    return (list(rows[0]), rows[1:]) if rows else ([], [])


def group_counts(sheet):
    """Counted from the Группа column when the sheet has one; an organisation
    filed under two groups counts in both, so the groups don't add up."""
    header, rows = sheet_rows(sheet)
    counts = collections.Counter()
    if 'Группа' in header:
        k = header.index('Группа')
        for r in rows:
            for g in str(r[k] or '').split(' / '):
                counts[g.strip()] += 1
    return len(rows), counts


def column_name(wb, *names):
    """The first of names found in the book's headers: the generated sheets say
    Карточка, the raw ones keep the export's maps_url."""
    for sheet in wb:
        header, _ = sheet_rows(sheet)
        for name in names:
            if name in header:
                return f'«{name}»'
    return f'«{names[0]}»'


def help_sheet(wb, title, contents, date, district_file=False):
    if 'Справка' in wb.sheetnames:
        wb.remove(wb['Справка'])
    ws = wb.create_sheet('Справка', 0)
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=15)
    ws['A2'] = f'Данные зафиксированы {date}. Источник организаций — Яндекс Карты, отделений партии — moscow.er.ru.'
    ws['A2'].font = NOTE
    ws.append([])
    ws.append(['КАКИЕ ОРГАНИЗАЦИИ В ФАЙЛЕ'])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    ws.append(['Лист', 'Организаций', 'Что там'])
    head = ws.max_row
    for row in contents:
        ws.append(row)
        ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True, vertical='top')
        if row[0] == 'ВСЕГО':
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
    for cell in ws[head]:
        cell.font, cell.fill = HEAD, FILL
    ws.append([])
    ws.append(['ГРУППЫ'])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    for name, text in (
        ('Общепит', 'кофейни, кафе, быстрое питание, рестораны, пекарни, пиццерии, столовые, бары'),
        ('Продуктовая розница', 'супермаркеты, магазины продуктов, мяса, овощей и фруктов, рыбы, '
                                'орехов, кулинарии, пива, чая'),
        ('Пищевое производство', 'производства продуктов, пищевые ингредиенты и специи, оптовые поставщики'),
        ('Религиозные учреждения', 'православные храмы, часовни, религиозные объединения, мечети, синагога. '
                                   'В питание не входят, даже если Яндекс приписал им заодно кофейню'),
        ('Единая Россия', 'районные местные отделения и окружное отделение ЮЗАО; ведутся вручную, '
                          'Яндекс Картами не собирались'),
    ):
        ws.append([name, None, text])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True, vertical='top')
    ws.append([])
    ws.append(['КАК РАБОТАТЬ С ФАЙЛОМ'])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    steps = [
        'Одна строка — одна организация (одна точка на карте). Одинаковые названия по разным адресам — '
        'это разные точки, их не удаляют.',
        'В строке заголовков включены фильтры: по ним отбирают нужные рубрики, строки с телефоном или сайтом.',
        f'Ссылка в колонке {column_name(wb, "Карточка", "maps_url")} открывает карточку на Яндекс Картах — '
        'по ней проверяют, работает ли организация и верен ли телефон.',
        f'Колонка {column_name(wb, "Сеть", "chain_flag")} отмечает точки одной сети. '
        'Их удобнее обрабатывать вместе: у сети часто общий телефон.',
        'Район определён по координатам точки на карте, а не по тому, каким запросом она найдена.',
    ]
    if not district_file:
        steps.insert(0, 'Листы районов пронумерованы 01–12 — так же, как отдельные файлы по районам.')
    for k, text in enumerate(steps, 1):
        ws.append([f'{k}.', None, text])
        ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True, vertical='top')
    for i, width in enumerate([34, 14, 90], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    return ws


def food_note(counts):
    parts = [f'{g.lower()} {counts[g]}' for g in FOOD_GROUPS if counts.get(g)]
    if counts.get('—'):
        parts.append(f'без рубрики {counts["—"]}')
    return ', '.join(parts)


def renumber_summary(ws, numbers):
    """Summary rows follow the district numbers instead of size."""
    found = [i for i in range(1, ws.max_row + 1) if plain(str(ws.cell(i, 1).value or '')) in numbers]
    if not found:
        return
    values = {plain(str(ws.cell(i, 1).value)): [c.value for c in ws[i]] for i in found}
    for i, name in zip(found, sorted(values, key=numbers.get)):
        for c, v in enumerate(values[name], 1):
            ws.cell(i, c).value = v
        ws.cell(i, 1).value = f'{numbers[name]:02d} {name}'


def finalize(out, per_district=None, date=None):
    date = date or datetime.date.today().strftime('%d.%m.%Y')
    files = {}
    if per_district:
        for path in sorted(glob.glob(os.path.join(per_district, '[0-9][0-9]_*.xlsx'))):
            stem = os.path.splitext(os.path.basename(path))[0]
            files[stem[3:].replace('_', ' ')] = (int(stem[:2]), path)

    district_stats = {}
    for district, (n, path) in files.items():
        book = load_workbook(path)
        total, counts = group_counts(book['Организации']) if 'Организации' in book.sheetnames else (0, {})
        rel = len(sheet_rows(book['Религия'])[1]) if 'Религия' in book.sheetnames else 0
        er = len(sheet_rows(book['Единая Россия'])[1]) if 'Единая Россия' in book.sheetnames else 0
        district_stats[district] = (total, counts)
        contents = [['Организации', total, 'питание: ' + (food_note(counts) or '—')]]
        if rel:
            contents.append(['Религия', rel, 'религиозные учреждения района'])
        if er:
            contents.append(['Единая Россия', er, 'отделение района и окружное отделение ЮЗАО'])
        help_sheet(book, f'{n:02d}. РАЙОН {district.upper()}', contents, date, district_file=True)
        book.save(path)

    wb = load_workbook(out)
    districts = [plain(n) for n in wb.sheetnames if plain(n) not in FIXED_SHEETS + ('Все организации', 'Сети')]
    numbers = {d: files[d][0] for d in districts if d in files}
    for d in sorted(d for d in districts if d not in numbers):
        numbers[d] = max(numbers.values(), default=0) + 1
    for sheet in wb:
        if plain(sheet.title) in numbers:
            sheet.title = f'{numbers[plain(sheet.title)]:02d} {plain(sheet.title)}'
    if 'Сводка' in wb.sheetnames:
        renumber_summary(wb['Сводка'], numbers)

    contents, food_total, mismatch = [], 0, []
    for name in ('Единая Россия', 'Религия'):
        if name in wb.sheetnames:
            what = ('отделения партии: районные и окружное ЮЗАО' if name == 'Единая Россия'
                    else 'религиозные учреждения всех районов')
            contents.append([name, len(sheet_rows(wb[name])[1]), what])
    for d in sorted(numbers, key=numbers.get):
        sheet = wb[f'{numbers[d]:02d} {d}']
        total, counts = group_counts(sheet)
        if not counts and d in district_stats:
            if district_stats[d][0] != total:
                mismatch.append(f'{d}: в общем файле {total}, в файле района {district_stats[d][0]}')
            counts = district_stats[d][1]
        food_total += total
        contents.append([sheet.title, total, 'питание: ' + (food_note(counts) or '—')])
    contents.append(['ВСЕГО', food_total + sum(r[1] for r in contents[:2] if r[0] in FIXED_SHEETS),
                     f'питание {food_total} + религия и отделения партии'])
    help_sheet(wb, 'РЕЕСТР ОРГАНИЗАЦИЙ ЮЗАО', contents, date)

    first = [n for n in FIXED_SHEETS if n in wb.sheetnames]
    ordered = first + [f'{numbers[d]:02d} {d}' for d in sorted(numbers, key=numbers.get)]
    ordered += [n for n in wb.sheetnames if n not in ordered]
    wb._sheets = [wb[n] for n in ordered]
    wb.active = 0
    for sheet in wb:
        sheet.sheet_view.tabSelected = sheet.title == 'Справка'
    wb.save(out)

    print(f'{out}: листы — ' + ', '.join(ordered))
    if files:
        print(f'{per_district}: справка добавлена в {len(files)} файлов')
    for line in mismatch:
        print('РАСХОЖДЕНИЕ', line)


def build(args):
    rows = read_final(args.final)
    religion = [r for r in rows if RELIGION in groups_of(r)]
    food = [r for r in rows if RELIGION not in groups_of(r)]
    order = lambda r: (GROUP_ORDER.get(label(r).split(' / ')[0], 9),
                       (r['categories'] or '').lower(), (r['title'] or '').lower())
    food = annotate(sorted(food, key=lambda r: (r['detected_district'],) + order(r)))
    religion = annotate(sorted(religion, key=lambda r: (r['detected_district'], (r['title'] or '').lower())))
    by_district = collections.OrderedDict()
    for row in food:
        by_district.setdefault(row['detected_district'], []).append(row)

    header, er = er_sheets(args.er) if args.er else (None, [])

    wb = Workbook()
    ws = wb.active
    ws.title = 'Сводка'
    ws['A1'] = 'РЕЕСТР ОРГАНИЗАЦИЙ'
    ws['A1'].font = Font(bold=True, size=15)
    ws['A2'] = 'Яндекс Карты. Дедупликация по place_id, район — по координатам.'
    ws['A2'].font = NOTE
    ws.append([]); ws.append([])
    ws.append(['ПИТАНИЕ — лист на каждый район, всё вместе на листе «Все организации»'])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    ws.append(['Район', 'Организаций', 'Общепит', 'Продуктовая розница', 'Пищевое производство',
               'Сетевых', 'Телефон есть', 'Сайт есть'])
    head_row = ws.max_row
    count = lambda v, g: sum(1 for r in v if g in groups_of(r))
    filled = lambda v, k: sum(1 for r in v if (r.get(k) or '').strip())
    line = lambda name, v: [name, len(v), count(v, 'Общепит'), count(v, 'Продуктовая розница'),
                            count(v, 'Пищевое производство'),
                            sum(1 for r in v if r.get('chain_flag') == 'CHAIN'),
                            filled(v, 'phone'), filled(v, 'website')]
    for district, v in sorted(by_district.items(), key=lambda x: -len(x[1])):
        ws.append(line(district, v))
    ws.append(line('ВСЕГО', food))
    total_row = ws.max_row
    ws.append([]); ws.append([])
    ws.append(['ОТДЕЛЬНО ОТ ПИТАНИЯ'])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    ws.append(['Раздел', 'Организаций', 'Лист', 'Комментарий'])
    head_row2 = ws.max_row
    ws.append(['Религиозные учреждения', len(religion), 'Религия',
               'Колонки «По адресу» и «Возможный дубль» — для ручной проверки'])
    if er:
        ws.append(['Отделения партии', len(er), 'Единая Россия', 'Из отдельного файла, Картами не собиралось'])
    for row_index in (head_row, head_row2):
        for cell in ws[row_index]:
            cell.font, cell.fill = HEAD, FILL
            cell.alignment = Alignment(wrap_text=True, vertical='center')
    for cell in ws[total_row]:
        cell.font = Font(bold=True)
    for i, width in enumerate([30, 14, 12, 22, 22, 12, 14, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    start = ws.max_row + 2
    for k, text in enumerate([
        'Строка с религиозной рубрикой не попадает в питание, даже если Яндекс приписал ей заодно кофейню.',
        '«Сеть» — организации с общим телефоном или названием. Для рассылки такие строки стоит схлопывать.',
        '«Сайт» и «Карточка» — кликабельные ссылки.',
        'Исходные данные со всеми служебными полями — в CSV-выгрузке расширения.',
    ]):
        ws.cell(start + k, 1, text).font = NOTE
    ws.freeze_panes = 'A6'

    working_sheet(wb, 'Все организации', food, district_col=True)
    for district, v in sorted(by_district.items()):
        working_sheet(wb, district, v)
    if religion:
        working_sheet(wb, 'Религия', religion, dup=True)

    sheet = wb.create_sheet('Сети')
    sheet.append(['Сеть (название)', 'Точек', 'Районы', 'Телефон'])
    chains = collections.defaultdict(list)
    for row in food:
        if row.get('chain_flag') == 'CHAIN':
            chains[(row['title'] or '').strip().lower()].append(row)
    for _, v in sorted(chains.items(), key=lambda x: -len(x[1])):
        phones = [r['phone'] for r in v if r['phone']]
        sheet.append([v[0]['title'], len(v), ', '.join(sorted({r['detected_district'] for r in v})),
                      collections.Counter(phones).most_common(1)[0][0] if phones else ''])
    for cell in sheet[1]:
        cell.font, cell.fill = HEAD, FILL
    for i, width in enumerate([34, 9, 58, 20], 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:D{sheet.max_row}'

    if er:
        sheet = wb.create_sheet('Единая Россия')
        sheet.append(header)
        for row in er:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font, cell.fill = HEAD, FILL
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        for i, width in enumerate([18, 26, 46, 30, 44, 22, 28, 16, 14, 30, 40], 1):
            sheet.column_dimensions[get_column_letter(i)].width = width
        sheet.freeze_panes = 'A2'

    wb.save(args.out)

    if args.per_district:
        os.makedirs(args.per_district, exist_ok=True)
        rel_by_district = collections.defaultdict(list)
        for row in religion:
            rel_by_district[row['detected_district']].append(row)
        er_by_district = collections.defaultdict(list)
        for row in er:
            er_by_district[(row[0] or '').strip()].append(row)
        for n, (district, v) in enumerate(sorted(by_district.items()), 1):
            book = Workbook()
            book.remove(book.active)
            working_sheet(book, 'Организации', v)
            if rel_by_district.get(district):
                working_sheet(book, 'Религия', rel_by_district[district], dup=True)
            own_er = er_by_district.get(district, []) + er_by_district.get('ЮЗАО', [])
            if own_er and header:
                sheet = book.create_sheet('Единая Россия')
                sheet.append(header)
                for row in own_er:
                    sheet.append(row)
                for cell in sheet[1]:
                    cell.font, cell.fill = HEAD, FILL
                    cell.alignment = Alignment(wrap_text=True, vertical='center')
                for i, width in enumerate([18, 26, 46, 30, 44, 22, 28, 16, 14, 30, 40], 1):
                    sheet.column_dimensions[get_column_letter(i)].width = width
                sheet.freeze_panes = 'A2'
            name = f"{n:02d}_{district.replace(' ', '_')}.xlsx"
            book.save(os.path.join(args.per_district, name))
        print(f'{args.per_district}: {len(by_district)} файлов по районам')

    if args.csv_dir:
        os.makedirs(args.csv_dir, exist_ok=True)
        fields = [f for f in rows[0] if not f.startswith('_')]
        def dump(path, data):
            with io.open(path, 'w', encoding='utf-8-sig', newline='') as f:
                writer = csv.DictWriter(f, fieldnames=fields, extrasaction='ignore', lineterminator='\r\n')
                writer.writeheader()
                writer.writerows(data)
        dump(os.path.join(args.csv_dir, 'FINAL_pitanie.csv'), food)
        if religion:
            dump(os.path.join(args.csv_dir, 'FINAL_religiya.csv'), religion)

    print(f'{args.out}: питание {len(food)}, религия {len(religion)}, сетей {len(chains)}'
          + (f', отделений партии {len(er)}' if er else ''))
    flagged = [r for r in religion if r['_dup']]
    if flagged:
        print(f'помечено как возможные дубли: {len(flagged)} строк (колонка «Возможный дубль»)')
    finalize(args.out, args.per_district, args.date)


def main():
    p = argparse.ArgumentParser(description='Рабочий отчёт из выгрузки FINAL_FILTERED')
    p.add_argument('--final', help='CSV выгрузки FINAL_FILTERED из расширения')
    p.add_argument('--finalize', action='store_true',
                   help='не собирать заново, а доработать готовые --out и --per-district: '
                        'справка, порядок листов, номера районов')
    p.add_argument('--date', help='дата фиксации данных для справки, ДД.ММ.ГГГГ (по умолчанию сегодня)')
    p.add_argument('--er', help='xlsx с отделениями партии (ведётся вручную, необязателен)')
    p.add_argument('--out', default='REESTR_otchet.xlsx', help='куда сохранить книгу')
    p.add_argument('--csv-dir', help='также выгрузить питание и религию отдельными CSV')
    p.add_argument('--per-district', help='папка для отдельных xlsx по каждому району')
    args = p.parse_args()
    if args.finalize:
        finalize(args.out, args.per_district, args.date)
    elif args.final:
        build(args)
    else:
        p.error('нужен --final (собрать отчёт) или --finalize (доработать готовый)')


if __name__ == '__main__':
    main()
